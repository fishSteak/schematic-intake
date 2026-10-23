# -*- coding: utf-8 -*-
"""Build boss-facing version comparison PPT from current copy."""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = r"K:\project\schematic-intake\docs\AI审图_电子_版本效果对照_v2.pptx"
FONT = "Microsoft YaHei"

HEAD = RGBColor(0x16, 0x3A, 0x6B)
ACCENT = RGBColor(0x1E, 0x4B, 0x8C)
TEXT = RGBColor(0x15, 0x20, 0x33)
MUTED = RGBColor(0x5A, 0x6A, 0x7C)
LINE = RGBColor(0xC5, 0xD0, 0xE0)
ALT = RGBColor(0xF5, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAGE = RGBColor(0xEE, 0xF1, 0xF6)
HI = RGBColor(0x0B, 0x6B, 0x3A)
NAV_OFF = RGBColor(0xD5, 0xDE, 0xE9)


def set_run(run, size: float, bold: bool = False, color=TEXT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(r_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", FONT)


def add_rect(slide, left, top, width, height, fill, line=None, line_w=Pt(1)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = line_w
    return sh


def add_textbox(slide, left, top, width, height, text, size=12, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def fill_cell_lines(cell, lines, *, sizes, bolds, colors, align):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        set_run(run, sizes[i], bolds[i], colors[i])


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_rect(slide, 0, 0, prs.slide_width, prs.slide_height, PAGE)

    # nav
    nav_y = Inches(0.28)
    nav_h = Inches(0.36)
    x = Inches(0.45)
    for label, on, w in [
        ("痛点 / 目标 / 进展", False, 1.9),
        ("设计思路", False, 1.35),
        ("效果 / 成果", True, 1.45),
    ]:
        add_rect(slide, x, nav_y, Inches(w), nav_h, HEAD if on else NAV_OFF)
        box = add_textbox(
            slide,
            x,
            nav_y + Inches(0.04),
            Inches(w),
            nav_h,
            label,
            size=11,
            bold=on,
            color=WHITE if on else MUTED,
            align=PP_ALIGN.CENTER,
        )
        x += Inches(w) - Inches(0.08)

    add_textbox(
        slide,
        Inches(0.45),
        Inches(0.78),
        Inches(12.4),
        Inches(0.3),
        "验证口径：机械成型机产品中心原理图样本  ·  检查清单约 85 项  ·  按检查大类统计准确率",
        size=12,
        color=TEXT,
    )
    add_textbox(
        slide,
        Inches(0.45),
        Inches(1.1),
        Inches(12.4),
        Inches(0.26),
        "一次审图看两边：业务交什么 · 我们检什么（样本板）",
        size=11,
        color=MUTED,
    )

    kpis = [
        ("6 类", "业务提供的材料", ["原理图 · BOM · 网表", "工程文件 · 工况 · 任务书"]),
        ("19 项", "当前优先检测项", ["全量清单约 85 项", "产品中心先落地这 19 项"]),
        ("94 个", "板上要检的器件", ["另约 70 条网络一并覆盖", "随板卡规模变化"]),
        ("44 份", "自动对照的规格书", ["按物料型号自搜入库", "审图时直接引用参数"]),
    ]
    kpi_y = Inches(1.38)
    kpi_h = Inches(1.42)
    kpi_gap = Inches(0.16)
    kpi_left = Inches(0.45)
    kpi_w = (Inches(12.4) - 3 * kpi_gap) / 4

    for i, (num, title, desc_lines) in enumerate(kpis):
        left = kpi_left + i * (kpi_w + kpi_gap)
        add_rect(slide, left, kpi_y, kpi_w, kpi_h, WHITE, LINE, Pt(1))
        add_rect(slide, left, kpi_y, kpi_w, Inches(0.055), ACCENT)
        add_textbox(slide, left + Inches(0.12), kpi_y + Inches(0.12), kpi_w - Inches(0.2), Inches(0.4), num, size=22, bold=True, color=ACCENT)
        add_textbox(slide, left + Inches(0.12), kpi_y + Inches(0.52), kpi_w - Inches(0.2), Inches(0.28), title, size=12, bold=True, color=TEXT)
        box = slide.shapes.add_textbox(left + Inches(0.12), kpi_y + Inches(0.8), kpi_w - Inches(0.2), Inches(0.55))
        tf = box.text_frame
        tf.word_wrap = True
        for li, line in enumerate(desc_lines):
            p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            set_run(run, 10, False, MUTED)

    # table
    shape = slide.shapes.add_table(4, 5, Inches(0.45), Inches(2.98), Inches(12.4), Inches(3.2))
    table = shape.table
    for i, w in enumerate([Inches(2.4), Inches(1.7), Inches(1.7), Inches(1.3), Inches(5.3)]):
        table.columns[i].width = w

    headers = [
        ["方案"],
        ["图纸规范与", "工程文件检查"],
        ["器件选型与", "物料可靠性"],
        ["其他类别"],
        ["做了哪些改动?"],
    ]
    for j, lines in enumerate(headers):
        cell = table.cell(0, j)
        fill_cell_lines(
            cell,
            lines,
            sizes=[11] * len(lines),
            bolds=[True] * len(lines),
            colors=[WHITE] * len(lines),
            align=PP_ALIGN.CENTER if j in (1, 2, 3) else PP_ALIGN.LEFT,
        )
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEAD
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    rows = [
        {
            "name": "v0.1  大模型直审",
            "name_color": TEXT,
            "acc": None,
            "upgrade": "以大模型按规则与器件手册直接审查，输出带证据的审图结论。",
            "merged": ["综合准确率 48%", "首轮样本 · 全项汇总"],
            "alt": False,
        },
        {
            "name": "v0.2  算法引擎 + 智能复核",
            "name_color": TEXT,
            "acc": ["88%", "46%", "探索中"],
            "upgrade": (
                "1. 增加电路解析器，构建原理图知识图谱，作为Agent核心上下文；"
                "2. 增加算法引擎，检测每个器件的每条通路"
                "（不同通路电压、电流不同），Agent智能复核为辅。"
            ),
            "merged": None,
            "alt": True,
        },
        {
            "name": "v0.3  Agent智能审图主链路",
            "name_color": HI,
            "acc": ["验证中", "验证中", "验证中"],
            "upgrade": (
                "1. 系统重构为混合智能架构，增加路由机制："
                "简单任务程序/算法处理，复杂任务交给Agent，兼顾性能与泛化；"
                "2. 转向Agent智能审图，算法保留作计算旁证与门禁；"
                "同类器件按任务分片并行审查。"
            ),
            "merged": None,
            "alt": False,
        },
    ]

    for i, row in enumerate(rows):
        r = i + 1
        bg = ALT if row["alt"] else WHITE

        cell0 = table.cell(r, 0)
        fill_cell_lines(
            cell0,
            [row["name"]],
            sizes=[11],
            bolds=[True],
            colors=[row["name_color"]],
            align=PP_ALIGN.LEFT,
        )
        cell0.fill.solid()
        cell0.fill.fore_color.rgb = bg
        cell0.vertical_anchor = MSO_ANCHOR.MIDDLE

        if row["merged"]:
            table.cell(r, 1).merge(table.cell(r, 3))
            cell = table.cell(r, 1)
            fill_cell_lines(
                cell,
                row["merged"],
                sizes=[13, 10],
                bolds=[True, False],
                colors=[TEXT, MUTED],
                align=PP_ALIGN.CENTER,
            )
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        else:
            for j, val in enumerate(row["acc"]):
                cell = table.cell(r, j + 1)
                big = val.endswith("%") and len(val) <= 4
                fill_cell_lines(
                    cell,
                    [val],
                    sizes=[14 if big else 11],
                    bolds=[big],
                    colors=[TEXT],
                    align=PP_ALIGN.CENTER,
                )
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE

        cell4 = table.cell(r, 4)
        fill_cell_lines(
            cell4,
            [row["upgrade"]],
            sizes=[10],
            bolds=[False],
            colors=[TEXT],
            align=PP_ALIGN.LEFT,
        )
        cell4.fill.solid()
        cell4.fill.fore_color.rgb = bg
        cell4.vertical_anchor = MSO_ANCHOR.MIDDLE

    foot_lines = [
        "左侧为业务交来的材料；右侧三项为系统实际要检的范围与深度（样本板一次完整审图；器件数随板卡变化）。",
        "其他类别（电源完整性、时钟/复位/时序、接口与抗干扰、可生产性等）仍在能力建设中。准确率：专家认可的检查结论占比。",
    ]
    box = slide.shapes.add_textbox(Inches(0.45), Inches(6.35), Inches(12.4), Inches(0.85))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(foot_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, 9, False, MUTED)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
