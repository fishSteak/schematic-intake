# -*- coding: utf-8 -*-
"""Build design-thinking PPT slide for electronic AI schematic review."""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = r"K:\project\schematic-intake\docs\AI审图_电子_设计思路.pptx"
FONT = "Microsoft YaHei"

HEAD = RGBColor(0x16, 0x3A, 0x6B)
TEXT = RGBColor(0x15, 0x20, 0x33)
MUTED = RGBColor(0x5A, 0x6A, 0x7C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, size: float, bold: bool = False, color=TEXT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(r_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", FONT)


def add_para(tf, text, *, size=12, bold=False, color=TEXT, space_before=0, space_after=4):
    p = tf.paragraphs[0] if not tf.paragraphs[0].runs else tf.add_paragraph()
    if not tf.paragraphs[0].runs and tf.paragraphs[0] is p:
        pass
    else:
        if tf.paragraphs[0].runs and p is tf.paragraphs[0]:
            p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return p


def write_block(tf, lines):
    """lines: list of (text, size, bold, color, space_before)"""
    first = True
    for text, size, bold, color, sb in lines:
        if first:
            p = tf.paragraphs[0]
            p.clear()
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(sb)
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = text
        set_run(run, size, bold, color)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # white bg
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.45))
    tf = title.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "2、设计思路（重要方案）"
    set_run(run, 22, True, HEAD)

    body = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(12.3), Inches(6.5))
    tf = body.text_frame
    tf.word_wrap = True

    blocks = [
        ("第一阶段：先验证「大模型能否直接审图」", True, HEAD, 10),
        (
            "调研大模型按清单与手册直接审图的能力。结论：缺少可复核的工程真值底座时，结论易漂、难回归，综合准确率约 48%。",
            False,
            TEXT,
            2,
        ),
        (
            "形成 v0.1 大模型直审方案：能出带证据报告，但不足以作为业务主判定链路；下一阶段必须补齐「工程文件 → 可计算上下文」。",
            False,
            TEXT,
            2,
        ),
        ("第二阶段：引入电路知识图谱 + 确定性算法引擎", True, HEAD, 12),
        (
            "研究 6 类工程材料如何统一建模，以及哪些检查可改由算法复现求解。结论：规范/工程一致性/部分硬阈值必须先解析为电路知识图谱，再以算法核验——类似机械侧「尺寸必须以 3D 为真值源」。",
            False,
            TEXT,
            2,
        ),
        (
            "形成 v0.2 算法引擎 + 智能复核：优先落地 19 项（全量约 85）；样本板约 94 器件 / 70 网络 / 44 份规格书；图纸规范类约 88%，选型可靠性约 46%。硬规则可算法化已打通，选型/应力类成为下一阶段重点。",
            False,
            TEXT,
            2,
        ),
        ("第三阶段：复杂判定交给 Agent，算法转为旁证与门禁", True, HEAD, 12),
        (
            "结论：开放性电气与选型推理由 Agent 主导；算法保留计算旁证、格式门禁与可回归验收。简单项走算法、复杂项走 Agent，并按器件分片并行。",
            False,
            TEXT,
            2,
        ),
        (
            "形成现行 v0.3 Agent 智能审图主链路（混合智能架构 + 路由）。准确率按大类重新标定中，待业务确认后回写效果页。",
            False,
            TEXT,
            2,
        ),
        ("误报漏报归因与下一步方向", True, HEAD, 12),
        (
            "偏差主要来自：拓扑/工况覆盖不完整、规格书参数缺口、检查口径与专家金样未完全对齐、算法项与 Agent 项边界仍需收紧。",
            False,
            TEXT,
            2,
        ),
        (
            "下一步先与业务对齐判定口径与抽检金样，再扩展其余检查大类。目标：优先类准确率达业务抽检要求，且关键结论均可复核、可回归。",
            False,
            TEXT,
            2,
        ),
    ]

    first = True
    for text, bold, color, sb in blocks:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(sb)
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = text
        set_run(run, 13 if bold else 12, bold, color)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
