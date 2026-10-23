# -*- coding: utf-8 -*-
"""Replace PPT slide 13 with P04-A03 case using real schematic PDF crop."""

from __future__ import annotations

import shutil
from pathlib import Path

import fitz
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC_CANDIDATES = [
    Path(r"\\nas.hanscnc.com\dev-userdata$\yangjr204840\Desktop\AI审图_v0.2-20261009.pptx"),
    Path(r"K:\project\schematic-intake\docs\AI审图_v0.2_p13_P04-A03.pptx"),
    Path(r"K:\project\schematic-intake\docs\_nas_AI审图_v0.2.pptx"),
]
OUT = Path(r"K:\project\schematic-intake\docs\AI审图_v0.2_p13_P04-A03.pptx")
NAS_OUT = Path(r"\\nas.hanscnc.com\dev-userdata$\yangjr204840\Desktop\AI审图_v0.2-20261009.pptx")
PDF = Path(r"K:\my-project\claude-code-p\0918全量材料\SCH_原版拆分-PDF重建-R1_2026-09-11(1).pdf")
CROP_DIR = Path(r"K:\project\schematic-intake\docs\_ppt_crops")

FONT = "微软雅黑"
RED = RGBColor(0xDF, 0x00, 0x3A)
FAIL = RGBColor(0xD0, 0x00, 0x28)
DARK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC8, 0xD0, 0xDC)
CARD = RGBColor(0xF7, 0xF9, 0xFC)
ACCENT = RGBColor(0x1E, 0x4B, 0x8C)


def set_run(run, size, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", FONT)


def write_text(shape, parts, *, align=PP_ALIGN.LEFT, size=12, bold=False, color=DARK):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(parts, str):
        parts = [(parts, bold, color, size)]
    for item in parts:
        if len(item) == 3:
            text, b, c = item
            sz = size
        else:
            text, b, c, sz = item
        run = p.add_run()
        run.text = text
        set_run(run, sz, b, c)


def add_rect(slide, left, top, width, height, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def add_round(slide, left, top, width, height, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def add_textbox(slide, left, top, width, height, parts, *, align=PP_ALIGN.LEFT, size=12, bold=False, color=DARK):
    box = slide.shapes.add_textbox(left, top, width, height)
    write_text(box, parts, align=align, size=size, bold=bold, color=color)
    return box


def delete_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def resolve_src() -> Path:
    for p in SRC_CANDIDATES:
        try:
            if p.exists():
                return p
        except OSError:
            continue
    raise FileNotFoundError("no source pptx")


def make_crop() -> Path:
    CROP_DIR.mkdir(exist_ok=True)
    out = CROP_DIR / "P04-A03_LEVEL2_ppt.png"
    doc = fitz.open(str(PDF))
    page = doc[5]  # PDF page 6 = P05_LEVELS
    rects = page.search_for("LEVEL2") + page.search_for("R26") + page.search_for("R27")
    for name in ("LEVEL1", "LEVEL3"):
        for r in page.search_for(name):
            cx = (rects[0].x0 + rects[0].x1) / 2
            cy = (rects[0].y0 + rects[0].y1) / 2
            if abs(r.x0 - cx) < 120 and abs(r.y0 - cy) < 120:
                rects.append(r)
    clip = fitz.Rect(
        min(r.x0 for r in rects) - 55,
        min(r.y0 for r in rects) - 45,
        max(r.x1 for r in rects) + 100,
        max(r.y1 for r in rects) + 55,
    ) & page.rect
    for r in page.search_for("LEVEL2"):
        a = page.add_rect_annot(fitz.Rect(r.x0 - 3, r.y0 - 3, r.x1 + 3, r.y1 + 3))
        a.set_colors(stroke=(0.82, 0, 0.23))
        a.set_border(width=2)
        a.update()
    pix = page.get_pixmap(matrix=fitz.Matrix(4, 4), clip=clip, alpha=False)
    pix.save(str(out))
    doc.close()
    return out


def main() -> None:
    crop = make_crop()
    src = resolve_src()
    print("src", src, "crop", crop)
    shutil.copy2(src, OUT)
    prs = Presentation(str(OUT))
    slide = prs.slides[12]

    for sh in list(slide.shapes):
        if sh.shape_type is not None and int(sh.shape_type) == 13 and sh.top < Emu(600000) and sh.left == 0:
            continue
        delete_shape(sh)

    add_rect(slide, Emu(0), Emu(107156), Emu(5393531), Emu(400050), RED)
    add_textbox(
        slide,
        Emu(114300),
        Emu(171450),
        Emu(5214938),
        Emu(300038),
        [
            ("实验研究阶段性成果", False, WHITE, 16.5),
            ("：", False, WHITE, 16.5),
            ("针对图纸文件种类的检查", True, WHITE, 16.5),
        ],
    )

    add_textbox(
        slide,
        Emu(400050),
        Emu(560000),
        Emu(8429625),
        Emu(280000),
        "AI 检索图面网络标号并统计同名出现次数，提前发现跨页漏标",
        size=17,
        bold=True,
        color=DARK,
    )
    add_textbox(
        slide,
        Emu(400050),
        Emu(840000),
        Emu(8429625),
        Emu(200000),
        "检查项 P04-A03：不同逻辑网络不得误用同名标号；同一逻辑网络跨页名称必须一致",
        size=11,
        color=MUTED,
    )

    # left: real PDF crop
    pic_left = Emu(350000)
    pic_top = Emu(1100000)
    pic_w = Emu(5200000)
    pic_h = Emu(2650000)
    add_round(slide, pic_left - Emu(20000), pic_top - Emu(20000), pic_w + Emu(40000), pic_h + Emu(40000), WHITE, LINE)
    slide.shapes.add_picture(str(crop), pic_left, pic_top, width=pic_w, height=pic_h)
    add_textbox(
        slide,
        pic_left,
        pic_top + pic_h + Emu(20000),
        pic_w,
        Emu(200000),
        "截取自原理图 PDF 第 6 页（P05_LEVELS）· 红框为网标 LEVEL2",
        size=10,
        color=MUTED,
    )

    # right: finding panel
    rx = Emu(5750000)
    add_round(slide, rx, Emu(1100000), Emu(3000000), Emu(2700000), CARD, LINE)
    add_textbox(slide, rx + Emu(120000), Emu(1180000), Emu(2760000), Emu(280000), "检索结果（全图）", size=13, bold=True, color=ACCENT)
    add_textbox(slide, rx + Emu(120000), Emu(1500000), Emu(2760000), Emu(350000), "网名：LEVEL2", size=14, bold=True, color=DARK)
    add_textbox(
        slide,
        rx + Emu(120000),
        Emu(1850000),
        Emu(2760000),
        Emu(700000),
        "同名网标出现次数：1 处\n仅见：第 6 页（本截图）\n其余页：未再检出「LEVEL2」",
        size=12,
        color=DARK,
    )
    add_textbox(
        slide,
        rx + Emu(120000),
        Emu(2650000),
        Emu(2760000),
        Emu(900000),
        [
            ("不通过：", True, FAIL, 13),
            ("按漏标标出\n", True, FAIL, 13),
        ],
    )
    add_textbox(
        slide,
        rx + Emu(120000),
        Emu(3100000),
        Emu(2760000),
        Emu(550000),
        "规则：同一逻辑网靠网标跨页连通时，图上同名标号应≥2 处。全图只有 1 处 → 判漏标。",
        size=11,
        color=MUTED,
    )

    # bottom 4 steps
    step_y_bar = Emu(4029075)
    title_y = Emu(4143375)
    l1_y = Emu(4414838)
    l2_y = Emu(4593431)
    cols = [Emu(400050), Emu(2543175), Emu(4686300), Emu(6829425)]
    steps = [
        ("1  AI抽取图面网标", "从原理图 PDF / 工程文件抽取", "本例定位到网标 LEVEL2"),
        ("2  AI全图统计同名", "按网名聚合，数出现次数", "LEVEL2 全图仅 1 处"),
        ("3  AI对照跨页规则", "跨页连通：同名标号应≥2 处", "其余页未检出第二处"),
        ("4  AI判定并提出改进", "不通过：LEVEL2 漏标", "在连通对端补标后复查"),
    ]
    for i, (t, a, b) in enumerate(steps):
        x = cols[i]
        add_rect(slide, x, step_y_bar, Emu(1914525), Emu(21431), RED)
        add_textbox(slide, x, title_y, Emu(2021681), Emu(242888), t, size=12, bold=True, color=DARK)
        add_textbox(slide, x, l1_y, Emu(2043113), Emu(192881), a, size=11, color=DARK)
        add_textbox(slide, x, l2_y, Emu(2043113), Emu(192881), b, size=11, color=DARK)

    add_textbox(
        slide,
        Emu(400050),
        Emu(4953000),
        Emu(6972300),
        Emu(161925),
        "说明：不是预先指定「某一页缺表/缺标」，而是全图检索后发现只有 1 处同名网标。导线仍可能画通，但跨页网标规范不满足，故按漏标提示。",
        size=9,
        color=MUTED,
    )
    add_textbox(slide, Emu(7496175), Emu(4972050), Emu(1381125), Emu(161925), "www.hanscnc.com", size=9, color=MUTED, align=PP_ALIGN.RIGHT)

    prs.save(str(OUT))
    print("saved", OUT)
    try:
        shutil.copy2(OUT, NAS_OUT)
        print("NAS updated")
    except Exception as e:
        print("NAS copy failed:", e)


if __name__ == "__main__":
    main()
